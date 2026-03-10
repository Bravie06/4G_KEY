import pandas as pd
import datetime
import os
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

def get_downloads_folder():
    return os.path.join(os.path.expanduser('~'), 'Downloads')

def generate_report(raw_file_path, template_file_path, event_name):
    # 1. Read raw data
    try:
        df = pd.read_excel(raw_file_path, sheet_name=0)
    except Exception as e:
        raise Exception(f"Failed to read raw file: {str(e)}")

    # 2. Identify entity column
    if 'Group' in df.columns:
        entity_col = 'Group'
    elif 'Managed Element' in df.columns:
        entity_col = 'Managed Element'
    elif 'Managed\xa0Element' in df.columns:
        entity_col = 'Managed\xa0Element'
    elif 'ENBFunction Name' in df.columns:
        entity_col = 'ENBFunction Name'
    else:
        raise Exception("Could not find a valid entity column (Group, Managed Element, or ENBFunction Name) in the raw data.")

    # 3. Filter the data to the last 10 unique hours of `Begin Time`
    df['Begin Time'] = pd.to_datetime(df['Begin Time'])
    unique_times = sorted(df['Begin Time'].dropna().unique())
    last_10_times = unique_times[-10:] if len(unique_times) >= 10 else unique_times
    df_filtered = df[df['Begin Time'].isin(last_10_times)].copy()

    # 4. Extract KPIs from Template dynamically
    ordered_kpis = []
    if template_file_path and os.path.exists(template_file_path):
        try:
            # The template provided stores data mapping in sheet 1
            template_df = pd.read_excel(template_file_path, sheet_name=1, nrows=50)
            first_col = template_df.iloc[:, 0].dropna()
            kpis_in_template = first_col[first_col.astype(str).str.contains('Average of|Sum of', na=False)].unique().tolist()

            for kpi in kpis_in_template:
                if str(kpi).startswith('Average of '):
                    kpi_raw = str(kpi).replace('Average of ', '')
                elif str(kpi).startswith('Sum of '):
                    kpi_raw = str(kpi).replace('Sum of ', '')
                else:
                    kpi_raw = str(kpi)

                ordered_kpis.append(kpi_raw)

        except Exception as e:
            raise Exception(f"Error reading template file: {str(e)}")

    if not ordered_kpis:
        # Fallback to the requested defaults if no template provided or empty
        ordered_kpis = [
            'ORA_4G_Cell Availability, excluding BLU_ZTE(%)',
            'ORA_4G_Total TRAFFIC(DL+UL)(GB)',
            'ORA_4G_ERAB_Setup_SR_new(%)',
            'ORA_4G_DL_User_Throughput_New(kbps)',
            'ORA_4G_LTE_Drop_Call_Rate_WO_VoLTE_New(%)',
            'ORA_4G_CALL_SETUP_SUCCESS_RATE_New(%)'
        ]

    for kpi in ordered_kpis:
        if kpi not in df_filtered.columns:
            raise Exception(f"KPI '{kpi}' required by the template is missing in the raw data columns.")

    # Multiply percentage-based KPI values by 100 if they are stored as fractions.
    for col in df_filtered.columns:
        if '(%)' in col and pd.api.types.is_numeric_dtype(df_filtered[col]):
            if df_filtered[col].max() <= 1.0:
                df_filtered[col] = df_filtered[col] * 100

    # 5. Pivot/group the data to calculate the required KPIs for each entity
    agg_funcs = {}
    for col in ordered_kpis:
        if 'TRAFFIC' in col.upper():
            agg_funcs[col] = 'sum'
        else:
            agg_funcs[col] = 'mean'

    grouped = df_filtered.groupby([entity_col, 'Begin Time'], as_index=False).agg(agg_funcs)
    entities = grouped[entity_col].unique()

    return _generate_excel(grouped, entities, last_10_times, entity_col, ordered_kpis, event_name)

def _generate_excel(grouped, entities, last_10_times, entity_col, ordered_kpis, event_name):
    now = datetime.datetime.now()
    date_str = now.strftime('%Y%m%d')
    time_str = now.strftime('%H%M%S')

    filename = f"KEA_4G_{event_name}_{date_str}_{time_str}.xlsx"
    downloads = get_downloads_folder()
    os.makedirs(downloads, exist_ok=True)
    out_path = os.path.join(downloads, filename)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    font_bold = Font(bold=True)
    border_thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')

    fill_green = PatternFill(start_color='00B050', end_color='00B050', fill_type='solid')
    fill_red = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
    fill_faded_red = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    fill_white = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

    ws.cell(row=2, column=1, value="Row Labels").font = font_bold
    ws.cell(row=2, column=1).border = border_thin

    for col_idx, ts in enumerate(last_10_times, start=2):
        cell = ws.cell(row=2, column=col_idx, value=ts.strftime('%Y-%m-%d %H:%M:%S'))
        cell.font = font_bold
        cell.border = border_thin
        cell.alignment = align_center

    current_row = 3

    kpi_order = []
    for kpi in ordered_kpis:
        if 'TRAFFIC' in kpi.upper():
            kpi_order.append(('Sum of ' + kpi, kpi))
        else:
            kpi_order.append(('Average of ' + kpi, kpi))

    for entity in entities:
        cell = ws.cell(row=current_row, column=1, value=entity)
        cell.font = font_bold
        cell.border = border_thin
        for col_idx in range(2, len(last_10_times) + 2):
            ws.cell(row=current_row, column=col_idx).border = border_thin
        current_row += 1

        entity_data = grouped[grouped[entity_col] == entity]

        for kpi_display_name, kpi_col_name in kpi_order:
            row_cell = ws.cell(row=current_row, column=1, value=kpi_display_name)
            row_cell.border = border_thin

            for col_idx, ts in enumerate(last_10_times, start=2):
                val_row = entity_data[entity_data['Begin Time'] == ts]
                val = val_row[kpi_col_name].values[0] if not val_row.empty else None

                cell = ws.cell(row=current_row, column=col_idx)
                cell.border = border_thin
                cell.alignment = align_center

                if val is not None and not pd.isna(val):
                    val = round(float(val), 2)
                    cell.value = val

                    if 'Availability' in kpi_col_name:
                        if val >= 99:
                            cell.fill = fill_green
                        else:
                            cell.fill = fill_faded_red
                    elif 'CALL_SETUP_SUCCESS_RATE' in kpi_col_name:
                        if val < 98.5:
                            cell.fill = fill_faded_red
                        else:
                            cell.fill = fill_green
                    elif 'Drop_Call_Rate' in kpi_col_name:
                        if val <= 0.5:
                            cell.fill = fill_green
                        else:
                            cell.fill = fill_red
                    else:
                        cell.fill = fill_white

            current_row += 1

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width

    wb.save(out_path)
    return out_path
