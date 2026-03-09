import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from datetime import datetime
import os

def process_data(raw_file_path, template_file_path, output_file_path):
    try:
        # Load raw data
        df_raw = pd.read_excel(raw_file_path)

        # Ensure 'Begin Time' is datetime
        df_raw['Begin Time'] = pd.to_datetime(df_raw['Begin Time'])

        # Get unique times and select the last 10 hours
        unique_times = sorted(df_raw['Begin Time'].unique())
        last_10_hours = unique_times[-10:] if len(unique_times) >= 10 else unique_times

        # Filter raw data for the last 10 hours
        df_filtered = df_raw[df_raw['Begin Time'].isin(last_10_hours)]

        # Get unique sites
        sites = df_filtered['Managed\xa0Element'].unique()

        # Parse template to dynamically get KPIs
        df_temp = pd.read_excel(template_file_path, header=None)

        # First site starts at row index 4 in template
        # The rows below the site name contain the KPIs
        kpi_labels = []
        for i in range(5, min(25, len(df_temp))):
            val = str(df_temp.iloc[i, 0])
            if pd.isna(df_temp.iloc[i, 0]) or val == 'nan':
                break
            # A site name might start with LIT_ or CTR_, check if it looks like a KPI
            if "Average" in val or "Sum" in val or "(%)" in val or "(GB)" in val or "(kbps)" in val:
                kpi_labels.append(val)
            else:
                break # Reached the next site

        if not kpi_labels:
            # Fallback
            kpi_labels = [
                "Average of ORA_4G_Cell Availability, excluding BLU_ZTE(%)",
                "Sum of ORA_4G_Total TRAFFIC(DL+UL)(GB)",
                "Average of ORA_4G_ERAB_Setup_SR_new(%)",
                "Average of ORA_4G_DL_User_Throughput_New(kbps)",
                "Average of ORA_4G_LTE_Drop_Call_Rate_WO_VoLTE_New(%)",
                "Average of ORA_4G_CALL_SETUP_SUCCESS_RATE_New(%)"
            ]

        # Define KPI mapping
        # Extract base column name from template KPI string
        kpi_mapping = []
        for label in kpi_labels:
            base_col = label.replace("Average of ", "").replace("Sum of ", "")
            kpi_mapping.append((label, base_col))

        # Colors for fonts and background fills
        green_font = Font(color="006100")
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

        red_font = Font(color="9C0006")
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        # Styles
        thin_border = Border(left=Side(style='thin'),
                             right=Side(style='thin'),
                             top=Side(style='thin'),
                             bottom=Side(style='thin'))

        bold_font = Font(bold=True)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        indent_align = Alignment(horizontal="left", vertical="center", indent=1)

        # Create a new workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Report"

        # Add headers as per template structure
        c_label = ws.cell(row=2, column=2, value="Column Labels")

        r_label = ws.cell(row=3, column=1, value="Row Labels")
        r_label.font = bold_font
        r_label.alignment = left_align
        r_label.border = thin_border

        for col_idx, time_val in enumerate(last_10_hours, start=2):
            # Format time as string
            cell = ws.cell(row=3, column=col_idx, value=pd.to_datetime(time_val).strftime("%Y-%m-%d %H:%M:%S"))
            cell.font = bold_font
            cell.alignment = center_align
            cell.border = thin_border

        current_row = 4

        for site in sites:
            # Site header
            cell = ws.cell(row=current_row, column=1, value=site)
            cell.font = bold_font
            cell.alignment = left_align
            cell.border = thin_border

            # Empty cells for site row across timestamps
            for col_idx in range(2, 2 + len(last_10_hours)):
                empty_cell = ws.cell(row=current_row, column=col_idx, value="")
                empty_cell.border = thin_border

            current_row += 1

            site_data = df_filtered[df_filtered['Managed\xa0Element'] == site]

            for kpi_label, raw_kpi in kpi_mapping:
                cell_kpi = ws.cell(row=current_row, column=1, value=kpi_label)
                cell_kpi.alignment = indent_align
                cell_kpi.border = thin_border

                for col_idx, time_val in enumerate(last_10_hours, start=2):
                    val_series = site_data[site_data['Begin Time'] == time_val][raw_kpi]
                    if not val_series.empty:
                        val = val_series.values[0]
                        cell = ws.cell(row=current_row, column=col_idx, value=val)

                        cell.border = thin_border
                        cell.alignment = center_align
                        if isinstance(val, (int, float)):
                            cell.number_format = "0.00"

                        # Apply coloring logic
                        if val is not None:
                            if "Availability" in kpi_label:
                                if val >= 99:
                                    cell.font = green_font
                                    cell.fill = green_fill
                                else:
                                    cell.font = red_font
                                    cell.fill = red_fill
                            elif "CALL_SETUP_SUCCESS_RATE" in kpi_label:
                                if val < 98.5:
                                    cell.font = red_font
                                    cell.fill = red_fill
                                else:
                                    cell.font = green_font
                                    cell.fill = green_fill
                            elif "Drop_Call_Rate" in kpi_label:
                                if val <= 0.5:
                                    cell.font = green_font
                                    cell.fill = green_fill
                                else:
                                    cell.font = red_font
                                    cell.fill = red_fill
                    else:
                        empty_cell = ws.cell(row=current_row, column=col_idx, value="")
                        empty_cell.border = thin_border

                current_row += 1

        # Adjust column widths
        ws.column_dimensions['A'].width = 75
        for col in range(2, 2 + len(last_10_hours)):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 25

        # Save output
        wb.save(output_file_path)
        return True, "Processing completed successfully!"
    except Exception as e:
        return False, f"An error occurred: {str(e)}"

if __name__ == "__main__":
    import tkinter as tk
    from tkinter import filedialog, messagebox

    def select_raw_file():
        file_path = filedialog.askopenfilename(
            title="Select Raw Data Excel File",
            filetypes=[("Excel files", "*.xlsx *.xls")]
        )
        if file_path:
            raw_entry.delete(0, tk.END)
            raw_entry.insert(0, file_path)

    def select_template_file():
        file_path = filedialog.askopenfilename(
            title="Select Template Excel File",
            filetypes=[("Excel files", "*.xlsx *.xls")]
        )
        if file_path:
            template_entry.delete(0, tk.END)
            template_entry.insert(0, file_path)

    def select_output_file():
        file_path = filedialog.asksaveasfilename(
            title="Save Output File As",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")]
        )
        if file_path:
            output_entry.delete(0, tk.END)
            output_entry.insert(0, file_path)

    def run_process():
        raw_path = raw_entry.get()
        template_path = template_entry.get()
        output_path = output_entry.get()

        if not raw_path or not template_path or not output_path:
            messagebox.showwarning("Missing Input", "Please select all required files and output location.")
            return

        # Disable button during processing
        run_btn.config(state=tk.DISABLED)
        root.update()

        success, msg = process_data(raw_path, template_path, output_path)

        run_btn.config(state=tk.NORMAL)

        if success:
            messagebox.showinfo("Success", msg)
        else:
            messagebox.showerror("Error", msg)

    root = tk.Tk()
    root.title("Excel Automation Tool")
    root.geometry("600x250")
    root.resizable(False, False)

    # Raw File
    tk.Label(root, text="Raw Data File:").grid(row=0, column=0, padx=10, pady=10, sticky="e")
    raw_entry = tk.Entry(root, width=50)
    raw_entry.grid(row=0, column=1, padx=10, pady=10)
    tk.Button(root, text="Browse...", command=select_raw_file).grid(row=0, column=2, padx=10, pady=10)

    # Template File
    tk.Label(root, text="Template File:").grid(row=1, column=0, padx=10, pady=10, sticky="e")
    template_entry = tk.Entry(root, width=50)
    template_entry.grid(row=1, column=1, padx=10, pady=10)
    tk.Button(root, text="Browse...", command=select_template_file).grid(row=1, column=2, padx=10, pady=10)

    # Output File
    tk.Label(root, text="Output File:").grid(row=2, column=0, padx=10, pady=10, sticky="e")
    output_entry = tk.Entry(root, width=50)
    output_entry.grid(row=2, column=1, padx=10, pady=10)
    tk.Button(root, text="Browse...", command=select_output_file).grid(row=2, column=2, padx=10, pady=10)

    # Run Button
    run_btn = tk.Button(root, text="Run Processing", command=run_process, bg="#4CAF50", fg="white", font=("Helvetica", 12, "bold"))
    run_btn.grid(row=3, column=1, pady=20)

    root.mainloop()
